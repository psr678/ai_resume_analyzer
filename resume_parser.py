"""
resume_parser.py

Extracts raw text from an uploaded resume file (PDF or DOCX).

Responsible AI note: this module does not store uploaded resume files
permanently. It only reads bytes passed in-memory (or an explicit path
for local testing) and returns text -- callers are responsible for not
writing the raw file to disk unless the user has explicitly agreed to it.
"""

import io
import os

from pypdf import PdfReader
from docx import Document

MAX_FILE_SIZE_MB = 5
ALLOWED_EXTENSIONS = {".pdf", ".docx"}


class ResumeParseError(Exception):
    """Raised when a resume file cannot be validated or parsed."""
    pass


def validate_file(filename: str, file_size_bytes: int) -> None:
    """Validate file extension and size before attempting to parse.

    Raises ResumeParseError with a human-readable message if invalid.
    """
    ext = os.path.splitext(filename)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ResumeParseError(
            f"Unsupported file type '{ext}'. Please upload a PDF or DOCX resume."
        )

    max_bytes = MAX_FILE_SIZE_MB * 1024 * 1024
    if file_size_bytes > max_bytes:
        raise ResumeParseError(
            f"File is too large ({file_size_bytes / (1024*1024):.1f} MB). "
            f"Please upload a resume under {MAX_FILE_SIZE_MB} MB."
        )


def extract_text_from_pdf(file_obj) -> str:
    """Extract text from every page of a PDF resume.

    file_obj can be a file path (str) or a file-like object (e.g. BytesIO
    from a Streamlit file_uploader), so this works both for the dashboard
    and for local/CLI testing against sample resumes on disk.
    """
    reader = PdfReader(file_obj)
    pages_text = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages_text.append(text)
    return "\n".join(pages_text)


def extract_text_from_docx(file_obj) -> str:
    """Extract text from every paragraph (and table cell) of a DOCX resume."""
    document = Document(file_obj)
    paragraphs = [p.text for p in document.paragraphs]

    # Many resumes use tables for layout (e.g. skills in a 2-column table),
    # so table cell text is included too, or that content would be silently lost.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.append(cell.text)

    return "\n".join(paragraphs)


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    """Main entry point: validate, then extract text based on file extension.

    filename is used only to determine the extension and for validation
    messages -- it is never used to write anything to disk.
    """
    validate_file(filename, len(file_bytes))
    ext = os.path.splitext(filename)[1].lower()
    file_obj = io.BytesIO(file_bytes)

    if ext == ".pdf":
        text = extract_text_from_pdf(file_obj)
    elif ext == ".docx":
        text = extract_text_from_docx(file_obj)
    else:
        # Defensive: validate_file should already have caught this.
        raise ResumeParseError(f"Unsupported file type '{ext}'.")

    if not text.strip():
        raise ResumeParseError(
            "No readable text could be extracted from this file. "
            "It may be a scanned/image-only resume, which this tool cannot read."
        )

    return text
