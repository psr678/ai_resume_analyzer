"""
Generates 3 sanitized sample resumes (no real personal information) as
both PDF and DOCX files, for testing resume_parser.py and the rest of
the pipeline. Resume A leans Data Analyst, B leans ML Engineer, C leans
NLP Engineer, matching the example test cases in the assignment PDF.
"""

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

resumes = {
    "resume_A_data_analyst": {
        "name": "Sample Candidate A",
        "summary": "Detail-oriented analyst with 2 years of experience turning raw "
                    "business data into clear dashboards and reports.",
        "skills": "Python, SQL, Excel, Pandas, Power BI, Data Visualization, Statistics, Git",
        "experience": [
            "Junior Data Analyst, Retail Insights Co. (2023-Present): Built weekly sales "
            "dashboards in Power BI, wrote SQL queries against a PostgreSQL warehouse, "
            "and used Pandas to clean and merge data from multiple sources.",
            "Data Analysis Intern, Campus Analytics Club (2022): Used Excel and Python "
            "to analyze survey data and presented statistical summaries to stakeholders.",
        ],
        "education": "B.Sc. in Statistics, State University (2023)",
        "projects": [
            "Sales Forecasting Dashboard: built an interactive Power BI dashboard on top "
            "of a SQL database to track regional sales trends.",
        ],
    },
    "resume_B_ml_engineer": {
        "name": "Sample Candidate B",
        "summary": "Machine learning engineer focused on building and deploying "
                    "production ML pipelines.",
        "skills": "Python, Machine Learning, scikit-learn, Pandas, NumPy, FastAPI, "
                  "Docker, MLflow, SQL, Git",
        "experience": [
            "ML Engineer, DataForge Labs (2022-Present): Trained and deployed scikit-learn "
            "models behind a FastAPI service, containerized with Docker, and tracked "
            "experiments with MLflow.",
            "Data Science Intern, University Research Lab (2021): Built machine learning "
            "models in Python using Pandas and scikit-learn for a research project.",
        ],
        "education": "B.Tech in Computer Science, Institute of Technology (2022)",
        "projects": [
            "Churn Prediction API: trained a scikit-learn classification model and served "
            "predictions through a FastAPI endpoint, deployed with Docker.",
        ],
    },
    "resume_C_nlp_engineer": {
        "name": "Sample Candidate C",
        "summary": "NLP-focused engineer experienced with modern transformer-based "
                    "language models.",
        "skills": "Python, NLP, Transformers, Hugging Face, spaCy, Deep Learning, "
                  "PyTorch, Machine Learning, Git",
        "experience": [
            "NLP Engineer, TextWorks AI (2023-Present): Fine-tuned Hugging Face "
            "transformer models for text classification, built preprocessing pipelines "
            "with spaCy, and trained models using PyTorch.",
            "Research Assistant, Language Technology Lab (2022): Worked on NLP research "
            "projects involving transformers and deep learning models for sentiment analysis.",
        ],
        "education": "M.Sc. in Computational Linguistics, Tech University (2023)",
        "projects": [
            "Sentiment Classifier: fine-tuned a Hugging Face transformer model with "
            "PyTorch and evaluated it on a labeled sentiment dataset.",
        ],
    },
}


def build_docx(key, data):
    doc = Document()
    doc.add_heading(data["name"], level=1)
    doc.add_paragraph(data["summary"])

    doc.add_heading("Skills", level=2)
    doc.add_paragraph(data["skills"])

    doc.add_heading("Experience", level=2)
    for exp in data["experience"]:
        doc.add_paragraph(exp, style="List Bullet")

    doc.add_heading("Projects", level=2)
    for proj in data["projects"]:
        doc.add_paragraph(proj, style="List Bullet")

    doc.add_heading("Education", level=2)
    doc.add_paragraph(data["education"])

    doc.save(f"sample_resumes/{key}.docx")


def build_pdf(key, data):
    styles = getSampleStyleSheet()
    story = [
        Paragraph(data["name"], styles["Title"]),
        Spacer(1, 8),
        Paragraph(data["summary"], styles["Normal"]),
        Spacer(1, 12),
        Paragraph("Skills", styles["Heading2"]),
        Paragraph(data["skills"], styles["Normal"]),
        Spacer(1, 12),
        Paragraph("Experience", styles["Heading2"]),
    ]
    for exp in data["experience"]:
        story.append(Paragraph("- " + exp, styles["Normal"]))
        story.append(Spacer(1, 4))

    story.append(Spacer(1, 8))
    story.append(Paragraph("Projects", styles["Heading2"]))
    for proj in data["projects"]:
        story.append(Paragraph("- " + proj, styles["Normal"]))
        story.append(Spacer(1, 4))

    story.append(Spacer(1, 8))
    story.append(Paragraph("Education", styles["Heading2"]))
    story.append(Paragraph(data["education"], styles["Normal"]))

    doc = SimpleDocTemplate(
        f"sample_resumes/{key}.pdf",
        pagesize=letter,
        topMargin=0.7 * inch, bottomMargin=0.7 * inch,
        leftMargin=0.7 * inch, rightMargin=0.7 * inch,
    )
    doc.build(story)


if __name__ == "__main__":
    import os
    os.makedirs("sample_resumes", exist_ok=True)
    for key, data in resumes.items():
        build_docx(key, data)
        build_pdf(key, data)
        print(f"Built {key}.docx and {key}.pdf")
